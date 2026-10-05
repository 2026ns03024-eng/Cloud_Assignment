#!/usr/bin/env python3
"""Build BITS WILP Cloud Computing Assignment 1 submission DOCX."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOT_DIR = Path(__file__).resolve().parent / "screenshots"
OUT = Path(__file__).resolve().parent / "Cloud_Computing_Assignment1_2026NS03024.docx"
BITS_ID = "2026NS03024"
SCRIPT_PATH = ROOT / "cloud_monitor.sh"

TERMINAL_SESSION = r"""2026NS03024> ./cloud_monitor.sh
Proceed with system scan? (y/n): y
{
  "timestamp": "2026-10-01T01:30:43+0530",
  "hostname": "MW659VFKXRD.cohesity.com",
  "metrics": {
    "disk_mount": "/",
    "disk_usage_percent": 5,
    "cpu_load_1m": 3.41,
    "cpu_load_5m": 3.66,
    "cpu_load_15m": 3.49
  },
  "log_analysis": {
    "log_file": "access.log",
    "total_requests": 1200,
    "count_2xx": 848,
    "count_4xx": 196,
    "count_5xx": 156,
    "error_samples": [ ... ]
  },
  "thresholds": { "max_disk_percent": 80, "error_limit": 5 },
  "alerts": {
    "disk_exceeded": false,
    "error_limit_exceeded": true,
    "webhook_sent": false
  }
}

2026NS03024> ./run_tests.sh
All automated tests passed."""

JSON_FIELDS = [
    ("timestamp", "ISO-style local time when the scan started"),
    ("hostname", "Host running the agent"),
    ("metrics.disk_mount", "Filesystem mount checked (root /)"),
    ("metrics.disk_usage_percent", "Used space percentage from df -P"),
    ("metrics.cpu_load_1m / 5m / 15m", "Load averages (proc/loadavg or sysctl on macOS)"),
    ("log_analysis.log_file", "Configured access log file name"),
    ("log_analysis.total_requests", "Log lines with a valid 3-digit HTTP status"),
    ("log_analysis.count_2xx / 4xx / 5xx", "HTTP status class counts"),
    ("log_analysis.error_samples", "Up to five sample 5xx rows (status, path, timestamp)"),
    ("thresholds.max_disk_percent", "MAX_DISK from ~/.monitor_conf"),
    ("thresholds.error_limit", "ERROR_LIMIT from ~/.monitor_conf"),
    ("alerts.disk_exceeded", "True when disk usage >= MAX_DISK"),
    ("alerts.error_limit_exceeded", "True when count_5xx > ERROR_LIMIT"),
    ("alerts.webhook_sent", "True when curl POST to WEBHOOK_URL succeeded during an alert"),
]

SCREENSHOTS = [
    (
        "01_guardrail_and_json.png",
        "Core (8 marks): Valid JSON report after system metrics and log analysis.",
        "Core guardrail (2 marks): y/n prompt before scan; user entered y.",
    ),
    (
        "02_monitor_conf.png",
        "Add-on #2 — Configuration file: thresholds read from ~/.monitor_conf.",
        None,
    ),
    (
        "03_monitor_history.png",
        "Add-on #3 — Audit logging: ~/.monitor_history run timestamp, 5xx count, duration.",
        None,
    ),
    (
        "04_log_archive.png",
        "Add-on #4 — Log archival: compressed access logs in ~/archive with date stamp.",
        None,
    ),
    (
        "05_webhook_payload.png",
        "Add-on #1 — Webhook alerting: JSON POST when error_limit_exceeded (receiver URL in config).",
        None,
    ),
]


def add_code_block(doc: Document, text: str) -> None:
    for line in text.splitlines():
        p = doc.add_paragraph(line)
        p.style = "No Spacing"
        for run in p.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(8)


def main() -> None:
    doc = Document()
    title = doc.add_heading("Cloud Computing EC1 — Assignment 1 Submission", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"BITS ID (custom shell prompt): {BITS_ID}")
    doc.add_paragraph(
        "Telemetry agent: cloud_monitor.sh — collects disk/CPU metrics, parses access.log "
        "for HTTP anomalies, prints structured JSON, optional webhook alert, audit history, "
        "and log rotation."
    )

    doc.add_heading("1. JSON data format explanation", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "JSON field"
    hdr[1].text = "Meaning"
    for field, meaning in JSON_FIELDS:
        row = table.add_row().cells
        row[0].text = field
        row[1].text = meaning
    doc.add_paragraph(
        "Webhook rule: when disk_exceeded or error_limit_exceeded is true, the agent POSTs "
        "the same JSON body to WEBHOOK_URL via curl."
    )

    doc.add_heading("2. Live terminal session (proof)", level=1)
    doc.add_paragraph(
        "Excerpt from interactive run with customized prompt and automated test suite:"
    )
    add_code_block(doc, TERMINAL_SESSION)

    doc.add_heading("3. Screenshot evidence (rubric)", level=1)
    for filename, caption1, caption2 in SCREENSHOTS:
        path = SHOT_DIR / filename
        if path.is_file():
            doc.add_paragraph(caption1, style="List Bullet")
            if caption2:
                doc.add_paragraph(caption2, style="List Bullet")
            doc.add_picture(str(path), width=Inches(6.2))
        else:
            doc.add_paragraph(f"[Missing image: {filename}]")

    doc.add_heading("4. Complete source code — cloud_monitor.sh", level=1)
    code = SCRIPT_PATH.read_text(encoding="utf-8")
    add_code_block(doc, code)

    doc.add_heading("5. Setup commands used", level=1)
    setup = f"""cp monitor_conf.sample ~/.monitor_conf
./generate_access_log.sh access.log 1200
export PS1='{BITS_ID}> '
./cloud_monitor.sh
./run_tests.sh
./capture_screenshots.sh"""
    add_code_block(doc, setup)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
